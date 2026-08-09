#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""
AI 旅遊行程整理器 MVP
- 本機 stdlib HTTP server
- 匯入本機資料夾 / 上傳檔案
- 抽取照片 EXIF 時間/GPS（Pillow 可用時）
- 抽取 TXT/MD/CSV/JSON 文字；PDF/DOCX/XLSX 以可選套件支援
- 建立每日時間軸、地點清單、Leaflet 地圖 HTML、Markdown 報告
"""
from __future__ import annotations

import datetime as _dt
import html
import json
import mimetypes
import os
import re
import shutil
import sys
import threading
import time
import traceback
import urllib.parse
import urllib.request
from dataclasses import dataclass, asdict
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple

ROOT = Path(__file__).resolve().parent
APP_ROOT = Path(sys.executable).resolve().parent if getattr(sys, "frozen", False) else ROOT
BUNDLE_ROOT = Path(getattr(sys, "_MEIPASS", ROOT))
DATA = APP_ROOT / "data"
TRIPS = DATA / "trips"
UPLOADS = DATA / "uploads"
CONFIG_PATH = DATA / "config.json"
STATE_PATH = DATA / "state.json"
STATIC_FILES = {"/": "index.html", "/index.html": "index.html", "/app.js": "app.js", "/styles.css": "styles.css"}
LOCK = threading.RLock()

IMAGE_EXT = {".jpg", ".jpeg", ".tif", ".tiff", ".heic", ".png", ".webp"}
VIDEO_EXT = {".mp4", ".mov", ".m4v", ".avi", ".mkv", ".webm"}
TEXT_EXT = {".txt", ".md", ".csv", ".json", ".yaml", ".yml", ".log"}
DOC_EXT = {".pdf", ".docx", ".xlsx"}
SUPPORTED_EXT = IMAGE_EXT | VIDEO_EXT | TEXT_EXT | DOC_EXT

DEFAULT_CONFIG = {
    "app_name": "AI 旅遊行程整理器",
    "default_trip": "demo-trip",
    "timezone": "Asia/Taipei",
    "geocode_enabled": True,
    "geocode_provider": "OpenStreetMap Nominatim",
    "max_text_chars_per_file": 12000,
    "copy_imported_files": True,
}


def ensure_dirs() -> None:
    for p in [DATA, TRIPS, UPLOADS]:
        p.mkdir(parents=True, exist_ok=True)
    if not CONFIG_PATH.exists():
        write_json(CONFIG_PATH, DEFAULT_CONFIG)
    if not STATE_PATH.exists():
        write_json(STATE_PATH, {"ok": True, "last_run": None, "events": []})


def read_json(path: Path, fallback: Any) -> Any:
    try:
        return json.loads(path.read_text(encoding="utf-8"))
    except Exception:
        return fallback


def write_json(path: Path, obj: Any) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    tmp = path.with_suffix(path.suffix + ".tmp")
    tmp.write_text(json.dumps(obj, ensure_ascii=False, indent=2), encoding="utf-8")
    tmp.replace(path)


def slugify(s: str) -> str:
    s = re.sub(r"[^\w\-\u4e00-\u9fff]+", "-", s.strip(), flags=re.UNICODE)
    s = re.sub(r"-+", "-", s).strip("-")
    return s or "trip"


def now_iso() -> str:
    return _dt.datetime.now().astimezone().isoformat(timespec="seconds")


def add_event(message: str, level: str = "info") -> None:
    with LOCK:
        st = read_json(STATE_PATH, {})
        events = st.get("events", [])
        events.insert(0, {"time": now_iso(), "level": level, "message": message})
        st["events"] = events[:80]
        write_json(STATE_PATH, st)


def safe_under(root: Path, target: Path) -> bool:
    try:
        target.resolve().relative_to(root.resolve())
        return True
    except Exception:
        return False


def parse_dt_from_text(text: str) -> Optional[str]:
    patterns = [
        r"(20\d{2})[/-](\d{1,2})[/-](\d{1,2})[ T_\-]?(\d{1,2})?:?(\d{2})?",
        r"(\d{4})(\d{2})(\d{2})[_-]?(\d{2})?(\d{2})?",
    ]
    for pat in patterns:
        m = re.search(pat, text)
        if not m:
            continue
        y, mo, d = int(m.group(1)), int(m.group(2)), int(m.group(3))
        hh = int(m.group(4) or 12)
        mm = int(m.group(5) or 0)
        try:
            return _dt.datetime(y, mo, d, hh, mm).isoformat(timespec="minutes")
        except ValueError:
            pass
    return None


def file_mtime_iso(path: Path) -> str:
    return _dt.datetime.fromtimestamp(path.stat().st_mtime).isoformat(timespec="minutes")


def gps_to_float(value: Any, ref: str) -> Optional[float]:
    try:
        def conv(x: Any) -> float:
            if isinstance(x, tuple):
                return float(x[0]) / float(x[1])
            if hasattr(x, "numerator") and hasattr(x, "denominator"):
                return float(x.numerator) / float(x.denominator)
            return float(x)
        deg, minute, sec = value
        out = conv(deg) + conv(minute) / 60.0 + conv(sec) / 3600.0
        if ref in ("S", "W"):
            out = -out
        return out
    except Exception:
        return None


def image_metadata(path: Path) -> Dict[str, Any]:
    meta = {"taken_at": None, "lat": None, "lng": None, "exif_available": False}
    try:
        from PIL import Image, ExifTags  # type: ignore
        with Image.open(path) as im:
            exif = im.getexif()
            if not exif:
                return meta
            meta["exif_available"] = True
            tagmap = {ExifTags.TAGS.get(k, k): v for k, v in exif.items()}
            dt = tagmap.get("DateTimeOriginal") or tagmap.get("DateTime")
            if dt:
                try:
                    meta["taken_at"] = _dt.datetime.strptime(str(dt), "%Y:%m:%d %H:%M:%S").isoformat(timespec="minutes")
                except Exception:
                    pass

            # Pillow may expose GPSInfo as an integer IFD offset in exif.items();
            # get_ifd(34853) is the reliable path for many Sony/phone JPEGs.
            gps_raw = None
            try:
                gps_raw = exif.get_ifd(34853)  # GPSInfo IFD
            except Exception:
                gps_raw = tagmap.get("GPSInfo")
            if gps_raw:
                gps = {ExifTags.GPSTAGS.get(k, k): v for k, v in gps_raw.items()}
                lat = gps_to_float(gps.get("GPSLatitude"), gps.get("GPSLatitudeRef", "N"))
                lng = gps_to_float(gps.get("GPSLongitude"), gps.get("GPSLongitudeRef", "E"))
                meta["lat"], meta["lng"] = lat, lng
    except Exception:
        # Pillow / EXIF parsing is optional.
        pass
    return meta


def extract_text(path: Path, max_chars: int) -> Tuple[str, str]:
    ext = path.suffix.lower()
    try:
        if ext in TEXT_EXT:
            return path.read_text(encoding="utf-8", errors="ignore")[:max_chars], "text"
        if ext == ".pdf":
            try:
                import pypdf  # type: ignore
                reader = pypdf.PdfReader(str(path))
                text = "\n".join((page.extract_text() or "") for page in reader.pages)
                return text[:max_chars], "pypdf"
            except Exception as e:
                return f"[PDF 文字抽取需安裝 pypdf；目前略過：{e}]", "pdf-unavailable"
        if ext == ".docx":
            try:
                import docx  # type: ignore
                d = docx.Document(str(path))
                text = "\n".join(p.text for p in d.paragraphs)
                return text[:max_chars], "python-docx"
            except Exception as e:
                return f"[DOCX 文字抽取需安裝 python-docx；目前略過：{e}]", "docx-unavailable"
        if ext == ".xlsx":
            try:
                import openpyxl  # type: ignore
                wb = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
                rows = []
                for ws in wb.worksheets:
                    rows.append(f"# Sheet: {ws.title}")
                    for row in ws.iter_rows(values_only=True):
                        vals = [str(v) for v in row if v is not None]
                        if vals:
                            rows.append(" | ".join(vals))
                return "\n".join(rows)[:max_chars], "openpyxl"
            except Exception as e:
                return f"[XLSX 文字抽取需安裝 openpyxl；目前略過：{e}]", "xlsx-unavailable"
    except Exception as e:
        return f"[文字抽取失敗：{e}]", "error"
    return "", "none"


def guess_title_from_text(path: Path, text: str) -> str:
    stem = path.stem.replace("_", " ").replace("-", " ").strip()
    lines = [x.strip() for x in text.splitlines() if x.strip()]
    if lines:
        first = re.sub(r"\s+", " ", lines[0])[:80]
        if 4 <= len(first) <= 80:
            return first
    return stem or path.name


def classify(path: Path, text: str) -> str:
    s = (path.name + "\n" + text[:2000]).lower()
    rules = [
        ("航班", ["flight", "航班", "起飛", "arrival", "departure", "boarding", "機票", "airport"]),
        ("飯店", ["hotel", "booking", "check-in", "check in", "住宿", "飯店", "旅館", "入住"]),
        ("餐廳", ["restaurant", "dinner", "lunch", "breakfast", "晚餐", "午餐", "早餐", "餐廳", "訂位"]),
        ("交通", ["train", "metro", "bus", "taxi", "jr", "地鐵", "電車", "巴士", "車票"]),
        ("景點", ["museum", "temple", "park", "tower", "神社", "寺", "公園", "博物館", "展望台", "景點"]),
        ("票券", ["ticket", "qr", "voucher", "門票", "票券", "憑證"]),
    ]
    for label, kws in rules:
        if any(k in s for k in kws):
            return label
    if path.suffix.lower() in IMAGE_EXT:
        return "照片"
    if path.suffix.lower() in VIDEO_EXT:
        return "影片"
    return "備註"


def extract_possible_place(text: str, fallback: str) -> str:
    patterns = [
        r"(?:地址|地點|Location|Address)[:：]\s*([^\n\r]{3,120})",
        r"(?:飯店|Hotel)[:：]\s*([^\n\r]{3,80})",
        r"(?:餐廳|Restaurant)[:：]\s*([^\n\r]{3,80})",
    ]
    for pat in patterns:
        m = re.search(pat, text, flags=re.I)
        if m:
            return m.group(1).strip()
    return fallback


def known_place_coordinates(query: str) -> Tuple[Optional[float], Optional[float], Optional[str]]:
    """Small safety net for common travel hubs where fuzzy geocoding can pick the wrong country."""
    q = (query or "").lower()
    known = [
        (["桃園國際機場", "桃園機場", "tpe", "taoyuan airport"], 25.07965, 121.23422, "桃園國際機場 Terminal 1, Taiwan"),
        (["松山機場", "tsa", "taipei songshan"], 25.06972, 121.55250, "台北松山機場, Taiwan"),
        (["成田", "nrt", "narita"], 35.77199, 140.39285, "成田國際機場, Japan"),
        (["羽田", "hnd", "haneda"], 35.54939, 139.77984, "東京羽田機場, Japan"),
        (["關西機場", "kix", "kansai airport"], 34.43472, 135.24417, "關西國際機場, Japan"),
    ]
    for keys, lat, lng, name in known:
        if any(k in q for k in keys):
            return lat, lng, name
    return None, None, None


def geocode(query: str) -> Tuple[Optional[float], Optional[float], Optional[str]]:
    if not query or len(query) < 3:
        return None, None, None
    known_lat, known_lng, known_name = known_place_coordinates(query)
    if known_lat is not None:
        return known_lat, known_lng, known_name
    try:
        url = "https://nominatim.openstreetmap.org/search?" + urllib.parse.urlencode({
            "format": "jsonv2", "limit": "1", "accept-language": "zh-TW,zh,en", "q": query
        })
        req = urllib.request.Request(url, headers={"User-Agent": "travel-ai-organizer-local/0.1"})
        with urllib.request.urlopen(req, timeout=8) as resp:
            data = json.loads(resp.read().decode("utf-8"))
        if data:
            return float(data[0]["lat"]), float(data[0]["lon"]), data[0].get("display_name")
    except Exception:
        return None, None, None
    return None, None, None


def reverse_geocode(lat: float, lng: float) -> Optional[str]:
    """Best-effort GPS -> place label, cached locally and rounded to avoid repeated calls."""
    cache_path = DATA / "geocode_cache.json"
    key = f"rev:{round(float(lat), 4)},{round(float(lng), 4)}"
    cache = read_json(cache_path, {})
    if key in cache:
        return cache[key]
    try:
        url = "https://nominatim.openstreetmap.org/reverse?" + urllib.parse.urlencode({
            "format": "jsonv2", "lat": f"{lat:.6f}", "lon": f"{lng:.6f}",
            "zoom": "18", "addressdetails": "1", "accept-language": "zh-TW,zh,en"
        })
        req = urllib.request.Request(url, headers={"User-Agent": "travel-ai-organizer-local/0.1"})
        with urllib.request.urlopen(req, timeout=8) as resp:
            data = json.loads(resp.read().decode("utf-8"))
        addr = data.get("address") or {}
        label = data.get("name") or addr.get("tourism") or addr.get("amenity") or addr.get("road") or data.get("display_name")
        display = data.get("display_name")
        if label and display and label != display:
            label = f"{label}｜{display}"
        if label:
            cache[key] = label
            write_json(cache_path, cache)
            time.sleep(1.0)
            return label
    except Exception:
        cache[key] = None
        write_json(cache_path, cache)
    return None


def thumbnail_for_image(path: Path, trip_dir: Path) -> Optional[str]:
    try:
        from PIL import Image, ImageOps  # type: ignore
        rel = path.relative_to(trip_dir)
        thumb_rel = Path("thumbs") / rel.with_suffix(".jpg")
        thumb_path = trip_dir / thumb_rel
        if thumb_path.exists() and thumb_path.stat().st_mtime >= path.stat().st_mtime:
            return str(thumb_rel).replace("\\", "/")
        thumb_path.parent.mkdir(parents=True, exist_ok=True)
        with Image.open(path) as im:
            im = ImageOps.exif_transpose(im)
            im.thumbnail((360, 360))
            if im.mode not in ("RGB", "L"):
                im = im.convert("RGB")
            im.save(thumb_path, "JPEG", quality=82, optimize=True)
        return str(thumb_rel).replace("\\", "/")
    except Exception:
        return None


def discover_files(folder: Path) -> List[Path]:
    out = []
    for p in folder.rglob("*"):
        if "thumbs" in p.parts:
            continue
        if p.is_file() and p.suffix.lower() in SUPPORTED_EXT:
            out.append(p)
    return sorted(out, key=lambda p: (p.stat().st_mtime, p.name.lower()))


def analyze_trip(trip_name: str, source_folder: Optional[str] = None) -> Dict[str, Any]:
    cfg = read_json(CONFIG_PATH, DEFAULT_CONFIG)
    slug = slugify(trip_name or cfg.get("default_trip", "trip"))
    trip_dir = TRIPS / slug
    raw_dir = trip_dir / "raw"
    extracted_dir = trip_dir / "extracted"
    thumbs_dir = trip_dir / "thumbs"
    for p in [trip_dir, raw_dir, extracted_dir, thumbs_dir]:
        p.mkdir(parents=True, exist_ok=True)

    source = Path(source_folder).expanduser() if source_folder else raw_dir
    if source_folder:
        source = source.resolve()
        if not source.exists() or not source.is_dir():
            raise ValueError(f"來源資料夾不存在：{source}")
        if cfg.get("copy_imported_files", True):
            for src in discover_files(source):
                rel = src.relative_to(source)
                dst = raw_dir / rel
                dst.parent.mkdir(parents=True, exist_ok=True)
                if not dst.exists() or src.stat().st_mtime > dst.stat().st_mtime:
                    shutil.copy2(src, dst)
            source = raw_dir

    files = discover_files(source)
    events: List[Dict[str, Any]] = []
    places: Dict[str, Dict[str, Any]] = {}
    max_chars = int(cfg.get("max_text_chars_per_file", 12000))
    reverse_budget = 0  # keep analysis fast; GPS labels are shown immediately, reverse lookup can be added as a later async enrichment
    reverse_cache = read_json(DATA / "geocode_cache.json", {})

    for idx, path in enumerate(files, 1):
        ext = path.suffix.lower()
        text = ""
        extraction = "metadata"
        image_meta = {}
        if ext in IMAGE_EXT:
            image_meta = image_metadata(path)
        if ext in TEXT_EXT or ext in DOC_EXT:
            text, extraction = extract_text(path, max_chars)
            (extracted_dir / (path.stem + ".txt")).write_text(text, encoding="utf-8")

        title = guess_title_from_text(path, text)
        dt = image_meta.get("taken_at") or parse_dt_from_text(path.name) or parse_dt_from_text(text) or file_mtime_iso(path)
        category = classify(path, text)
        lat, lng = image_meta.get("lat"), image_meta.get("lng")
        geocoded_name = None
        place_name = extract_possible_place(text, title)
        if lat is not None and lng is not None:
            rev_key = f"rev:{round(float(lat), 4)},{round(float(lng), 4)}"
            if rev_key in reverse_cache:
                geocoded_name = reverse_cache.get(rev_key) or f"GPS {float(lat):.5f}, {float(lng):.5f}"
            elif reverse_budget > 0:
                geocoded_name = reverse_geocode(float(lat), float(lng)) or f"GPS {float(lat):.5f}, {float(lng):.5f}"
                reverse_budget -= 1
                reverse_cache[rev_key] = geocoded_name
            else:
                geocoded_name = f"GPS {float(lat):.5f}, {float(lng):.5f}"
            if category == "照片":
                title = path.stem
        elif cfg.get("geocode_enabled") and category not in {"照片", "影片", "備註"}:
            lat, lng, geocoded_name = geocode(place_name)
            time.sleep(1.0)  # Nominatim usage policy friendly
        else:
            if category in {"照片", "影片"}:
                place_name = "未定位"
        event_id = f"e{idx:04d}"
        rel_path = str(path.relative_to(trip_dir)) if safe_under(trip_dir, path) else str(path)
        rel_url = rel_path.replace("\\", "/")
        thumb_rel = rel_url if ext in IMAGE_EXT else None
        event = {
            "id": event_id,
            "title": title,
            "datetime": dt,
            "date": dt[:10] if dt else "未知日期",
            "time": dt[11:16] if dt and len(dt) >= 16 else "--:--",
            "category": category,
            "place": geocoded_name or place_name,
            "lat": lat,
            "lng": lng,
            "source_file": rel_path,
            "source_type": ext.lstrip("."),
            "media_url": f"/asset/{urllib.parse.quote(slug)}/{urllib.parse.quote(rel_url)}",
            "thumbnail_url": f"/thumb/{urllib.parse.quote(slug)}/{urllib.parse.quote(thumb_rel)}" if thumb_rel else None,
            "extractor": extraction,
            "summary": make_summary(category, title, text, path),
        }
        events.append(event)
        if lat is not None and lng is not None:
            key = f"{round(float(lat), 5)},{round(float(lng), 5)}"
            places.setdefault(key, {"name": event["place"], "lat": lat, "lng": lng, "events": []})["events"].append(event_id)

    events.sort(key=lambda e: (e.get("datetime") or "9999", e.get("title") or ""))
    timeline = {
        "trip_name": trip_name,
        "slug": slug,
        "generated_at": now_iso(),
        "source_folder": str(source),
        "event_count": len(events),
        "place_count": len(places),
        "events": events,
        "days": build_days(events),
    }
    write_json(trip_dir / "timeline.json", timeline)
    write_json(trip_dir / "places.json", list(places.values()))
    (trip_dir / "report.md").write_text(render_markdown(timeline), encoding="utf-8")
    (trip_dir / "map.html").write_text(render_map_html(timeline, list(places.values())), encoding="utf-8")

    with LOCK:
        st = read_json(STATE_PATH, {})
        st.update({"last_run": now_iso(), "last_trip": slug, "last_event_count": len(events), "ok": True})
        write_json(STATE_PATH, st)
    add_event(f"完成分析：{trip_name}，事件 {len(events)} 筆、地點 {len(places)} 個")
    return timeline


def make_summary(category: str, title: str, text: str, path: Path) -> str:
    if text:
        clean = re.sub(r"\s+", " ", text).strip()
        if clean and not clean.startswith("["):
            return clean[:180] + ("…" if len(clean) > 180 else "")
    return f"從 {path.name} 建立的{category}事件。"


def build_days(events: List[Dict[str, Any]]) -> List[Dict[str, Any]]:
    grouped: Dict[str, List[Dict[str, Any]]] = {}
    for e in events:
        grouped.setdefault(e.get("date") or "未知日期", []).append(e)
    return [{"date": d, "events": grouped[d]} for d in sorted(grouped)]


def render_markdown(tl: Dict[str, Any]) -> str:
    lines = [f"# {tl['trip_name']} AI 行程整理", "", f"產生時間：{tl['generated_at']}", "", f"事件數：{tl['event_count']}｜地點數：{tl['place_count']}", ""]
    for day in tl.get("days", []):
        lines += [f"## {day['date']}", ""]
        for e in day["events"]:
            loc = f"｜{e['place']}" if e.get("place") else ""
            coord = f"｜{e['lat']:.5f}, {e['lng']:.5f}" if e.get("lat") is not None and e.get("lng") is not None else ""
            lines.append(f"- **{e['time']}** [{e['category']}] {e['title']}{loc}{coord}")
            lines.append(f"  - 摘要：{e['summary']}")
            lines.append(f"  - 來源：`{e['source_file']}`")
        lines.append("")
    return "\n".join(lines)


def render_map_html(tl: Dict[str, Any], places: List[Dict[str, Any]]) -> str:
    points = [e for e in tl.get("events", []) if e.get("lat") is not None and e.get("lng") is not None]
    payload = json.dumps({"timeline": tl, "points": points, "places": places}, ensure_ascii=False)
    return f"""<!doctype html><html lang=\"zh-Hant\"><head><meta charset=\"utf-8\"><meta name=\"viewport\" content=\"width=device-width, initial-scale=1\"><title>{html.escape(tl['trip_name'])} 地圖</title><link rel=\"stylesheet\" href=\"https://unpkg.com/leaflet@1.9.4/dist/leaflet.css\"><style>body{{margin:0;font-family:system-ui,'Noto Sans TC',sans-serif;background:#07111f;color:#eef}}#map{{height:100vh}}.panel{{position:absolute;z-index:999;top:16px;left:16px;width:320px;max-height:80vh;overflow:auto;background:rgba(8,18,34,.9);border:1px solid #24476f;border-radius:16px;padding:14px;box-shadow:0 20px 60px #0008}}.item{{border-top:1px solid #274360;padding:8px 0}}small{{color:#9fc3ee}}</style></head><body><div id=\"map\"></div><div class=\"panel\"><h2>{html.escape(tl['trip_name'])}</h2><p>{tl['event_count']} 筆事件｜{tl['place_count']} 個地點</p><div id=\"list\"></div></div><script src=\"https://unpkg.com/leaflet@1.9.4/dist/leaflet.js\"></script><script>const DATA={payload};const map=L.map('map');L.tileLayer('https://{{s}}.tile.openstreetmap.org/{{z}}/{{x}}/{{y}}.png',{{maxZoom:19,attribution:'&copy; OpenStreetMap'}}).addTo(map);const pts=DATA.points;const bounds=[];pts.forEach((e,i)=>{{const m=L.marker([e.lat,e.lng]).addTo(map);m.bindPopup(`<b>${{e.time}} ${{e.title}}</b><br>${{e.category}}<br>${{e.place||''}}<br><small>${{e.summary||''}}</small>`);bounds.push([e.lat,e.lng]);}});if(bounds.length){{map.fitBounds(bounds,{{padding:[50,50]}})}}else{{map.setView([25.033,121.5654],12)}}document.getElementById('list').innerHTML=DATA.timeline.days.map(d=>`<h3>${{d.date}}</h3>`+d.events.map(e=>`<div class=item><b>${{e.time}}</b> [${{e.category}}] ${{e.title}}<br><small>${{e.place||''}}</small></div>`).join('')).join('');</script></body></html>"""


class Handler(BaseHTTPRequestHandler):
    server_version = "TravelAIOrganizer/0.1"

    def log_message(self, fmt: str, *args: Any) -> None:
        sys.stderr.write("[%s] %s\n" % (self.log_date_time_string(), fmt % args))

    def send_json(self, obj: Any, status: int = 200) -> None:
        data = json.dumps(obj, ensure_ascii=False, indent=2).encode("utf-8")
        self.send_response(status)
        self.send_header("Content-Type", "application/json; charset=utf-8")
        self.send_header("Content-Length", str(len(data)))
        self.end_headers()
        self.wfile.write(data)

    def read_body_json(self) -> Dict[str, Any]:
        n = int(self.headers.get("Content-Length", "0") or 0)
        raw = self.rfile.read(n) if n else b"{}"
        return json.loads(raw.decode("utf-8") or "{}")

    def do_GET(self) -> None:
        try:
            parsed = urllib.parse.urlparse(self.path)
            path = parsed.path
            if path == "/api/status":
                return self.send_json(self.api_status())
            if path == "/api/trips":
                return self.send_json({"ok": True, "trips": self.list_trips()})
            if path.startswith("/api/trip/"):
                slug = urllib.parse.unquote(path.split("/")[-1])
                tpath = TRIPS / slug / "timeline.json"
                if not tpath.exists():
                    return self.send_json({"ok": False, "error": "trip not found"}, 404)
                return self.send_json({"ok": True, "timeline": read_json(tpath, {})})
            if path.startswith("/thumb/"):
                parts = path[len("/thumb/"):].split("/", 1)
                if len(parts) != 2:
                    return self.send_error(404)
                slug = urllib.parse.unquote(parts[0])
                rel = urllib.parse.unquote(parts[1])
                trip_dir = TRIPS / slug
                src = (trip_dir / rel).resolve()
                if not safe_under(trip_dir, src) or not src.exists() or not src.is_file():
                    return self.send_error(404)
                thumb_rel = thumbnail_for_image(src, trip_dir)
                if not thumb_rel:
                    return self.send_error(404)
                target = (trip_dir / thumb_rel).resolve()
                data = target.read_bytes()
                self.send_response(200)
                self.send_header("Content-Type", "image/jpeg")
                self.send_header("Content-Length", str(len(data)))
                self.send_header("Cache-Control", "public, max-age=86400")
                self.end_headers()
                self.wfile.write(data)
                return
            if path.startswith("/asset/"):
                parts = path[len("/asset/"):].split("/", 1)
                if len(parts) != 2:
                    return self.send_error(404)
                slug = urllib.parse.unquote(parts[0])
                rel = urllib.parse.unquote(parts[1])
                target = (TRIPS / slug / rel).resolve()
                if not safe_under(TRIPS / slug, target) or not target.exists() or not target.is_file():
                    return self.send_error(404)
                ctype = mimetypes.guess_type(str(target))[0] or "application/octet-stream"
                data = target.read_bytes()
                self.send_response(200)
                self.send_header("Content-Type", ctype)
                self.send_header("Content-Length", str(len(data)))
                self.send_header("Cache-Control", "public, max-age=86400")
                self.end_headers()
                self.wfile.write(data)
                return
            if path.startswith("/download/"):
                rel = urllib.parse.unquote(path[len("/download/"):])
                target = (TRIPS / rel).resolve()
                if not safe_under(TRIPS, target) or not target.exists() or not target.is_file():
                    return self.send_error(404)
                ctype = mimetypes.guess_type(str(target))[0] or "application/octet-stream"
                data = target.read_bytes()
                self.send_response(200)
                self.send_header("Content-Type", ctype)
                self.send_header("Content-Length", str(len(data)))
                self.send_header("Content-Disposition", f"attachment; filename={target.name}")
                self.end_headers()
                self.wfile.write(data)
                return
            if path in STATIC_FILES:
                return self.serve_file(BUNDLE_ROOT / STATIC_FILES[path])
            return self.send_error(404)
        except Exception as e:
            traceback.print_exc()
            return self.send_json({"ok": False, "error": str(e)}, 500)

    def do_POST(self) -> None:
        try:
            path = urllib.parse.urlparse(self.path).path
            if path == "/api/config":
                body = self.read_body_json()
                cfg = read_json(CONFIG_PATH, DEFAULT_CONFIG)
                cfg.update({k: v for k, v in body.items() if k in DEFAULT_CONFIG})
                write_json(CONFIG_PATH, cfg)
                add_event("已更新設定")
                return self.send_json({"ok": True, "config": cfg})
            if path == "/api/analyze":
                body = self.read_body_json()
                trip_name = body.get("trip_name") or read_json(CONFIG_PATH, DEFAULT_CONFIG).get("default_trip")
                source_folder = body.get("source_folder") or None
                timeline = analyze_trip(trip_name, source_folder)
                return self.send_json({"ok": True, "timeline": timeline})
            return self.send_json({"ok": False, "error": "unknown endpoint"}, 404)
        except Exception as e:
            traceback.print_exc()
            add_event(f"錯誤：{e}", "error")
            return self.send_json({"ok": False, "error": str(e)}, 500)

    def serve_file(self, path: Path) -> None:
        if not path.exists():
            return self.send_error(404)
        ctype = mimetypes.guess_type(str(path))[0] or "text/html; charset=utf-8"
        data = path.read_bytes()
        self.send_response(200)
        self.send_header("Content-Type", ctype)
        self.send_header("Content-Length", str(len(data)))
        self.end_headers()
        self.wfile.write(data)

    def api_status(self) -> Dict[str, Any]:
        cfg = read_json(CONFIG_PATH, DEFAULT_CONFIG)
        st = read_json(STATE_PATH, {})
        optional = {}
        for mod in ["PIL", "pypdf", "docx", "openpyxl"]:
            try:
                __import__(mod)
                optional[mod] = True
            except Exception:
                optional[mod] = False
        return {"ok": True, "config": cfg, "state": st, "optional_modules": optional, "data_dir": str(DATA), "trips": self.list_trips()}

    def list_trips(self) -> List[Dict[str, Any]]:
        out = []
        for p in sorted(TRIPS.iterdir() if TRIPS.exists() else [], key=lambda x: x.name.lower()):
            if not p.is_dir():
                continue
            tl = read_json(p / "timeline.json", {}) if (p / "timeline.json").exists() else {}
            out.append({
                "slug": p.name,
                "name": tl.get("trip_name", p.name),
                "event_count": tl.get("event_count", 0),
                "place_count": tl.get("place_count", 0),
                "generated_at": tl.get("generated_at"),
                "downloads": {
                    "report_md": f"/download/{p.name}/report.md",
                    "map_html": f"/download/{p.name}/map.html",
                    "timeline_json": f"/download/{p.name}/timeline.json",
                }
            })
        return out


def create_demo_if_empty() -> None:
    demo = TRIPS / "demo-trip" / "raw"
    demo.mkdir(parents=True, exist_ok=True)
    sample = demo / "2026-04-12_0730_桃園機場集合.txt"
    if not sample.exists():
        sample.write_text("地點：桃園國際機場 T1\n航班：TPE → NRT\n備註：集合、辦理登機。", encoding="utf-8")
    sample2 = demo / "2026-04-12_1900_新宿晚餐.txt"
    if not sample2.exists():
        sample2.write_text("餐廳：新宿燒肉店\n地址：Shinjuku City, Tokyo, Japan\n備註：第一天抵達東京後的晚餐。", encoding="utf-8")


def main() -> None:
    ensure_dirs()
    create_demo_if_empty()
    if not (TRIPS / "demo-trip" / "timeline.json").exists():
        analyze_trip("demo-trip", None)
    port = int(os.environ.get("PORT", "8765"))
    host = os.environ.get("HOST", "127.0.0.1")
    httpd = ThreadingHTTPServer((host, port), Handler)
    print(f"AI 旅遊行程整理器啟動：http://{host}:{port}/", flush=True)
    httpd.serve_forever()


if __name__ == "__main__":
    main()
